from fastapi import APIRouter, HTTPException, Body
from fastapi.responses import StreamingResponse
from typing import Dict, Any, List, Optional
import io
from datetime import datetime
from collections import OrderedDict
import urllib.parse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/export", tags=["Rapor Dışa Aktarma & Excel"])

def build_mutabakat_sheet(ws, title: str, subtitle: str, data: List[Dict[str, Any]], now_str: str):
    """Tek bir Excel sayfasına (worksheet) stilize Mutabakat tablosu çizer."""
    ws.views.sheetView[0].showGridLines = True
    if not data or len(data) == 0:
        ws.cell(row=1, column=1, value="Seçili kriterlerde veri bulunamadı.")
        return

    # Kategori sütununu normal tablodan çıkarıp şerit başlık ve grup toplamı olarak kullanıyoruz
    columns = [c for c in data[0].keys() if c != "Kategori"]
    col_count = len(columns)

    # Renk paleti
    c_banner_top = "0A1C15"      # Gece Zümrüt
    c_banner_sub = "0F291E"      # Koyu Orman Yeşili
    c_header_bg = "133829"       # Kolon Başlığı Yeşili
    c_header_border = "10B981"   # Canlı Zümrüt Bordür
    c_cat_banner = "19382B"      # Kategori Başlık Şeridi
    c_cat_border = "34D399"      # Kategori Şerit Çizgisi
    c_subtotal_bg = "ECFDF5"     # Kategori Toplam Zemin (Açık Nane)
    c_subtotal_border = "059669" # Kategori Toplam Alt Çizgisi
    c_grand_bg = "0A2419"        # Genel Toplam Zemin
    c_grand_border = "10B981"    # Genel Toplam Bordürü
    c_row_alt = "F8FAF9"         # Açık Nane Alternatif Satır
    c_border_light = "E2E8F0"    # İnce Gri Hücre Çizgisi

    font_title = Font(name="Segoe UI", size=13, bold=True, color="4ADE80")
    font_sub = Font(name="Segoe UI", size=9, italic=True, color="D1FAE5")
    font_header = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    font_cat_banner = Font(name="Segoe UI", size=10.5, bold=True, color="6EE7B7")
    font_subtotal_label = Font(name="Segoe UI", size=9.5, bold=True, color="064E3B")
    font_subtotal_val = Font(name="Segoe UI", size=9.5, bold=True, color="064E3B")
    font_grand_label = Font(name="Segoe UI", size=10.5, bold=True, color="FDE047")
    font_grand_val = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9.5, color="0F172A")
    font_data_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_code = Font(name="Segoe UI", size=9, color="475569")
    font_unit = Font(name="Segoe UI", size=9, color="64748B")
    font_red = Font(name="Segoe UI", size=9.5, bold=True, color="DC2626")
    font_green = Font(name="Segoe UI", size=9.5, bold=True, color="16A34A")
    font_red_grand = Font(name="Segoe UI", size=10.5, bold=True, color="FECACA")
    font_green_grand = Font(name="Segoe UI", size=10.5, bold=True, color="A7F3D0")

    fill_banner_top = PatternFill(start_color=c_banner_top, end_color=c_banner_top, fill_type="solid")
    fill_banner_sub = PatternFill(start_color=c_banner_sub, end_color=c_banner_sub, fill_type="solid")
    fill_header = PatternFill(start_color=c_header_bg, end_color=c_header_bg, fill_type="solid")
    fill_cat_banner = PatternFill(start_color=c_cat_banner, end_color=c_cat_banner, fill_type="solid")
    fill_subtotal = PatternFill(start_color=c_subtotal_bg, end_color=c_subtotal_bg, fill_type="solid")
    fill_grand = PatternFill(start_color=c_grand_bg, end_color=c_grand_bg, fill_type="solid")
    fill_alt = PatternFill(start_color=c_row_alt, end_color=c_row_alt, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_red_soft = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_green_soft = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_border_light),
        bottom=Side(style='thin', color=c_border_light)
    )
    header_border = Border(
        left=Side(style='thin', color=c_header_border),
        right=Side(style='thin', color=c_header_border),
        top=Side(style='medium', color=c_header_border),
        bottom=Side(style='medium', color=c_header_border)
    )
    cat_border = Border(
        left=Side(style='thin', color=c_cat_border),
        right=Side(style='thin', color=c_cat_border),
        top=Side(style='thin', color=c_cat_border),
        bottom=Side(style='thin', color=c_cat_border)
    )
    subtotal_border = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_cat_border),
        bottom=Side(style='medium', color=c_subtotal_border)
    )
    grand_border = Border(
        left=Side(style='thin', color=c_grand_border),
        right=Side(style='thin', color=c_grand_border),
        top=Side(style='medium', color=c_grand_border),
        bottom=Side(style='double', color=c_grand_border)
    )

    # 1. SATIR: ÜST BAŞLIK BANNERI
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    cell_t = ws.cell(row=1, column=1, value=f"👑 RAPOR ASİSTAN  |  {title}")
    cell_t.font = font_title
    cell_t.fill = fill_banner_top
    cell_t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, col_count + 1):
        ws.cell(row=1, column=c).fill = fill_banner_top

    # 2. SATIR: ALT BİLGİ & TARİH
    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    cell_s = ws.cell(row=2, column=1, value=f"📊 {subtitle}  •  İndirilme: {now_str}")
    cell_s.font = font_sub
    cell_s.fill = fill_banner_sub
    cell_s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, col_count + 1):
        ws.cell(row=2, column=c).fill = fill_banner_sub

    # 3. SATIR: BOŞLUK
    ws.row_dimensions[3].height = 6

    # 4. SATIR: TABLO KOLON BAŞLIKLARI
    ws.row_dimensions[4].height = 28
    for col_idx, col_name in enumerate(columns, start=1):
        cell_h = ws.cell(row=4, column=col_idx, value=str(col_name).upper())
        cell_h.font = font_header
        cell_h.fill = fill_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = header_border

    # Kategorilere göre grupla
    cat_groups = OrderedDict()
    for r in data:
        c_name = (r.get("Kategori") or r.get("Kategori (Özelkod6)") or r.get("Özelkod6") or r.get("Kategori Adı") or "").strip()
        if not c_name or c_name == "None":
            c_name = "DİĞER"
        if c_name not in cat_groups:
            cat_groups[c_name] = []
        cat_groups[c_name].append(r)

    current_row = 5

    for cat_name, items in cat_groups.items():
        # A. KATEGORİ BAŞLIK ŞERİDİ
        ws.row_dimensions[current_row].height = 24
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_count)
        cell_cat = ws.cell(row=current_row, column=1, value=f"📁  {cat_name.upper()}  ({len(items)} Kalem Ürün)")
        cell_cat.font = font_cat_banner
        cell_cat.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(1, col_count + 1):
            c_cell = ws.cell(row=current_row, column=c)
            c_cell.fill = fill_cat_banner
            c_cell.border = cat_border
        current_row += 1

        # B. ÜRÜN SATIRLARI
        for row_idx, item in enumerate(items):
            ws.row_dimensions[current_row].height = 20
            is_alt = (row_idx % 2 == 1)
            base_fill = fill_alt if is_alt else fill_white

            for col_idx, col_name in enumerate(columns, start=1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.border = thin_border
                cell.fill = base_fill

                val = item.get(col_name)

                if col_name == "Stok Kodu":
                    cell.value = str(val or "")
                    cell.font = font_code
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                elif col_name == "Ürün Adı":
                    cell.value = str(val or "")
                    cell.font = font_data_bold
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                elif col_name == "Birim":
                    cell.value = str(val or "ADET")
                    cell.font = font_unit
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name == "Hata %":
                    if val is None or val == "-":
                        cell.value = "-"
                        cell.font = font_unit
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        try:
                            num_ratio = float(val) / 100.0
                            cell.value = num_ratio
                            cell.number_format = '+0.0%;-0.0%;0.0%'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            if num_ratio < 0:
                                cell.font = font_red
                                cell.fill = fill_red_soft
                            elif num_ratio > 0:
                                cell.font = font_green
                                cell.fill = fill_green_soft
                            else:
                                cell.font = font_unit
                        except Exception:
                            cell.value = str(val)
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name == "Fark":
                    if val is None or val == "-":
                        cell.value = "-"
                        cell.font = font_unit
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        try:
                            num_fark = float(val)
                            cell.value = num_fark
                            cell.number_format = '+#,##0.00;-#,##0.00;0.00'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            if num_fark < 0:
                                cell.font = font_red
                                cell.fill = fill_red_soft
                            elif num_fark > 0:
                                cell.font = font_green
                                cell.fill = fill_green_soft
                            else:
                                cell.font = font_unit
                        except Exception:
                            cell.value = str(val)
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    # Sayısal alanlar (Önceki Sayım, Gelenler, Çıkanlar, Satış, Olması Gereken vb.)
                    if val is None or val == "-":
                        cell.value = 0
                        cell.number_format = '#,##0'
                        cell.font = font_unit
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        try:
                            num_val = float(val)
                            cell.value = num_val
                            cell.number_format = '#,##0.00' if (num_val % 1 != 0) else '#,##0'
                            cell.font = font_data
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        except Exception:
                            cell.value = str(val)
                            cell.font = font_data
                            cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        # C. KATEGORİ TOPLAMI SATIRI
        ws.row_dimensions[current_row].height = 22
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        cell_st_label = ws.cell(row=current_row, column=1, value=f"📊 {cat_name.upper()} TOPLAMI ({len(items)} Kalem)")
        cell_st_label.font = font_subtotal_label
        cell_st_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        cell_st_unit = ws.cell(row=current_row, column=3, value="-")
        cell_st_unit.font = font_subtotal_label
        cell_st_unit.alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, col_count + 1):
            ws.cell(row=current_row, column=c).fill = fill_subtotal
            ws.cell(row=current_row, column=c).border = subtotal_border

        for col_idx, col_name in enumerate(columns[3:], start=4):
            cell_sub = ws.cell(row=current_row, column=col_idx)
            if col_name == "Hata %":
                cat_fark = sum(float(it.get("Fark", 0) or 0) for it in items)
                cat_giris = sum(float(it.get("Toplam Giriş", 0) or it.get("Şubeye Gelenler", 0) or 0) for it in items)
                cat_base = cat_giris if cat_giris > 0 else sum(float(it.get("Toplam Çıkış", 0) or 0) for it in items)
                if cat_base <= 0:
                    cat_base = abs(sum(float(it.get("Olması Gereken", 0) or 0) for it in items))
                cat_ratio = (cat_fark / cat_base) if cat_base > 0 else 0.0
                cell_sub.value = cat_ratio
                cell_sub.number_format = '+0.0%;-0.0%;0.0%'
                cell_sub.alignment = Alignment(horizontal="right", vertical="center")
                if cat_ratio < 0:
                    cell_sub.font = font_red
                    cell_sub.fill = fill_red_soft
                elif cat_ratio > 0:
                    cell_sub.font = font_green
                    cell_sub.fill = fill_green_soft
                else:
                    cell_sub.font = font_subtotal_val
            elif col_name == "Fark":
                cat_fark = sum(float(it.get("Fark", 0) or 0) for it in items)
                cell_sub.value = cat_fark
                cell_sub.number_format = '+#,##0.00;-#,##0.00;0.00'
                cell_sub.alignment = Alignment(horizontal="right", vertical="center")
                if cat_fark < 0:
                    cell_sub.font = font_red
                    cell_sub.fill = fill_red_soft
                elif cat_fark > 0:
                    cell_sub.font = font_green
                    cell_sub.fill = fill_green_soft
                else:
                    cell_sub.font = font_subtotal_val
            else:
                cat_sum = sum(float(it.get(col_name, 0) or 0) for it in items)
                cell_sub.value = cat_sum
                cell_sub.number_format = '#,##0.00' if (cat_sum % 1 != 0) else '#,##0'
                cell_sub.font = font_subtotal_val
                cell_sub.alignment = Alignment(horizontal="right", vertical="center")

        current_row += 1
        ws.row_dimensions[current_row].height = 6
        current_row += 1

    # D. GENEL TOPLAM SATIRI (EN ALT)
    ws.row_dimensions[current_row].height = 26
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    cell_gt_label = ws.cell(row=current_row, column=1, value=f"👑 GENEL TOPLAM ({len(data)} Kalem Ürün)")
    cell_gt_label.font = font_grand_label
    cell_gt_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    cell_gt_unit = ws.cell(row=current_row, column=3, value="-")
    cell_gt_unit.font = font_grand_label
    cell_gt_unit.alignment = Alignment(horizontal="center", vertical="center")

    for c in range(1, col_count + 1):
        ws.cell(row=current_row, column=c).fill = fill_grand
        ws.cell(row=current_row, column=c).border = grand_border

    for col_idx, col_name in enumerate(columns[3:], start=4):
        cell_gt = ws.cell(row=current_row, column=col_idx)
        if col_name == "Hata %":
            grand_fark = sum(float(it.get("Fark", 0) or 0) for it in data)
            grand_giris = sum(float(it.get("Toplam Giriş", 0) or it.get("Şubeye Gelenler", 0) or 0) for it in data)
            grand_base = grand_giris if grand_giris > 0 else sum(float(it.get("Toplam Çıkış", 0) or 0) for it in data)
            if grand_base <= 0:
                grand_base = abs(sum(float(it.get("Olması Gereken", 0) or 0) for it in data))
            grand_ratio = (grand_fark / grand_base) if grand_base > 0 else 0.0
            cell_gt.value = grand_ratio
            cell_gt.number_format = '+0.0%;-0.0%;0.0%'
            cell_gt.alignment = Alignment(horizontal="right", vertical="center")
            if grand_ratio < 0:
                cell_gt.font = font_red_grand
            elif grand_ratio > 0:
                cell_gt.font = font_green_grand
            else:
                cell_gt.font = font_grand_val
        elif col_name == "Fark":
            grand_fark = sum(float(it.get("Fark", 0) or 0) for it in data)
            cell_gt.value = grand_fark
            cell_gt.number_format = '+#,##0.00;-#,##0.00;0.00'
            cell_gt.alignment = Alignment(horizontal="right", vertical="center")
            if grand_fark < 0:
                cell_gt.font = font_red_grand
            elif grand_fark > 0:
                cell_gt.font = font_green_grand
            else:
                cell_gt.font = font_grand_val
        else:
            grand_sum = sum(float(it.get(col_name, 0) or 0) for it in data)
            cell_gt.value = grand_sum
            cell_gt.number_format = '#,##0.00' if (grand_sum % 1 != 0) else '#,##0'
            cell_gt.font = font_grand_val
            cell_gt.alignment = Alignment(horizontal="right", vertical="center")

    # Kolon genişlikleri (Auto-fit)
    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        if col_name == "Stok Kodu":
            ws.column_dimensions[col_letter].width = 16
        elif col_name == "Ürün Adı":
            ws.column_dimensions[col_letter].width = 38
        elif col_name == "Birim":
            ws.column_dimensions[col_letter].width = 10
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)


def build_standard_sheet(ws, title: str, subtitle: str, data: List[Dict[str, Any]], now_str: str):
    """Standart liste ve tablo raporları için profesyonel Excel sayfası çizer."""
    ws.views.sheetView[0].showGridLines = True
    columns = list(data[0].keys())
    col_count = len(columns)

    c_banner_top = "0D1F18"
    c_banner_sub = "132A22"
    c_header_bg = "1A382E"
    c_header_border = "2CBE56"
    c_row_alt = "F4F9F6"
    c_border_light = "E2E8F0"

    font_title = Font(name="Segoe UI", size=14, bold=True, color="4ADE80")
    font_sub = Font(name="Segoe UI", size=9.5, italic=True, color="D1FAE5")
    font_header = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="0F172A")
    font_bold_num = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_red = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    font_green = Font(name="Segoe UI", size=10, bold=True, color="16A34A")

    fill_banner_top = PatternFill(start_color=c_banner_top, end_color=c_banner_top, fill_type="solid")
    fill_banner_sub = PatternFill(start_color=c_banner_sub, end_color=c_banner_sub, fill_type="solid")
    fill_header = PatternFill(start_color=c_header_bg, end_color=c_header_bg, fill_type="solid")
    fill_alt = PatternFill(start_color=c_row_alt, end_color=c_row_alt, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_red_soft = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_green_soft = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_border_light),
        bottom=Side(style='thin', color=c_border_light)
    )
    header_border = Border(
        left=Side(style='thin', color=c_header_border),
        right=Side(style='thin', color=c_header_border),
        top=Side(style='medium', color=c_header_border),
        bottom=Side(style='medium', color=c_header_border)
    )

    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    cell_t = ws.cell(row=1, column=1, value=f"👑 RAPOR ASİSTAN  |  {title}")
    cell_t.font = font_title
    cell_t.fill = fill_banner_top
    cell_t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col_idx in range(1, col_count + 1):
        ws.cell(row=1, column=col_idx).fill = fill_banner_top

    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    cell_s = ws.cell(row=2, column=1, value=f"📊 {subtitle}  •  İndirilme: {now_str}  •  Kurumsal Yönetim Raporu")
    cell_s.font = font_sub
    cell_s.fill = fill_banner_sub
    cell_s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col_idx in range(1, col_count + 1):
        ws.cell(row=2, column=col_idx).fill = fill_banner_sub

    ws.row_dimensions[3].height = 8

    ws.row_dimensions[4].height = 26
    for col_idx, col_name in enumerate(columns, start=1):
        cell_h = ws.cell(row=4, column=col_idx, value=str(col_name).upper())
        cell_h.font = font_header
        cell_h.fill = fill_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = header_border

    for row_idx, row_data in enumerate(data, start=5):
        ws.row_dimensions[row_idx].height = 20
        is_alt = (row_idx % 2 == 1)
        base_fill = fill_alt if is_alt else fill_white

        for col_idx, col_name in enumerate(columns, start=1):
            val = row_data.get(col_name)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.fill = base_fill

            col_lower = str(col_name).lower()
            is_currency = any(w in col_lower for w in ["ciro", "tutar", "fiyat", "bedel", "değer"])
            is_fark = "fark" in col_lower

            if isinstance(val, (int, float)):
                cell.value = val
                if is_currency:
                    cell.number_format = '₺ #,##0.00;[Red]-₺ #,##0.00;"-"'
                    cell.font = font_bold_num
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif is_fark:
                    cell.number_format = '+#,##0.00;-#,##0.00;0.00'
                    if val < 0:
                        cell.font = font_red
                        cell.fill = fill_red_soft
                    elif val > 0:
                        cell.font = font_green
                        cell.fill = fill_green_soft
                    else:
                        cell.font = font_data
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    if val < 0:
                        cell.font = font_red
                        cell.number_format = '#,##0.00'
                    else:
                        cell.font = font_data
                        cell.number_format = '#,##0.00' if isinstance(val, float) else '#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            elif val is None or val == "-":
                cell.value = "-"
                cell.font = font_data
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.value = str(val)
                cell.font = font_data
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for r_idx in range(5, min(5 + len(data), 100)):
            val_str = str(ws.cell(row=r_idx, column=col_idx).value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 5, 42))

    ws.freeze_panes = "A5"


def build_ai_studio_sheet(ws, title: str, subtitle: str, data: List[Dict[str, Any]], now_str: str):
    """
    AI Studio için dinamik sütunlu Excel sayfası.
    """
    ws.views.sheetView[0].showGridLines = True
    if not data or len(data) == 0:
        ws.cell(row=1, column=1, value="Dışa aktarılacak veri bulunamadı.")
        return

    columns = list(data[0].keys())
    col_count = len(columns)

    # Renk paleti
    c_banner_top = "064E3B"      # Deep emerald
    c_banner_sub = "0F766E"      # Teal emerald
    c_header_bg = "134E4A"       # Dark teal
    c_header_border = "10B981"   # Emerald border
    c_cat_bg = "065F46"          # Emerald category strip
    c_cat_border = "34D399"      # Light emerald
    c_subtotal_bg = "F0FDF4"     # Light green tint
    c_grand_bg = "064E3B"        # Deep emerald grand total
    c_row_alt = "F8FAFC"         # Clean slate alt
    c_border_light = "E2E8F0"

    font_title = Font(name="Segoe UI", size=13, bold=True, color="ECFDF5")
    font_sub = Font(name="Segoe UI", size=9, italic=True, color="D1FAE5")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_cat = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    font_subtotal = Font(name="Segoe UI", size=9.5, bold=True, color="065F46")
    font_grand = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_data_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_red = Font(name="Segoe UI", size=9.5, bold=True, color="DC2626")
    font_green = Font(name="Segoe UI", size=9.5, bold=True, color="16A34A")

    fill_banner_top = PatternFill(start_color=c_banner_top, end_color=c_banner_top, fill_type="solid")
    fill_banner_sub = PatternFill(start_color=c_banner_sub, end_color=c_banner_sub, fill_type="solid")
    fill_header = PatternFill(start_color=c_header_bg, end_color=c_header_bg, fill_type="solid")
    fill_cat = PatternFill(start_color=c_cat_bg, end_color=c_cat_bg, fill_type="solid")
    fill_subtotal = PatternFill(start_color=c_subtotal_bg, end_color=c_subtotal_bg, fill_type="solid")
    fill_grand = PatternFill(start_color=c_grand_bg, end_color=c_grand_bg, fill_type="solid")
    fill_alt = PatternFill(start_color=c_row_alt, end_color=c_row_alt, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_red_soft = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_green_soft = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_border_light),
        bottom=Side(style='thin', color=c_border_light)
    )
    header_border = Border(
        left=Side(style='thin', color=c_header_border),
        right=Side(style='thin', color=c_header_border),
        top=Side(style='medium', color=c_header_border),
        bottom=Side(style='medium', color=c_header_border)
    )
    cat_border = Border(
        left=Side(style='thin', color=c_cat_border),
        right=Side(style='thin', color=c_cat_border),
        top=Side(style='thin', color=c_cat_border),
        bottom=Side(style='thin', color=c_cat_border)
    )
    subtotal_border = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_cat_border),
        bottom=Side(style='medium', color=c_cat_border)
    )
    grand_border = Border(
        left=Side(style='thin', color="064E3B"),
        right=Side(style='thin', color="064E3B"),
        top=Side(style='medium', color="10B981"),
        bottom=Side(style='double', color="10B981")
    )

    # 1. Üst Başlık Bannerı
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    cell_t = ws.cell(row=1, column=1, value=f"👑 RAPOR ASİSTAN  |  {title}")
    cell_t.font = font_title
    cell_t.fill = fill_banner_top
    cell_t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, col_count + 1):
        ws.cell(row=1, column=c).fill = fill_banner_top

    # 2. Alt Başlık & Zaman
    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    cell_s = ws.cell(row=2, column=1, value=f"📊 {subtitle}  •  İndirilme: {now_str}")
    cell_s.font = font_sub
    cell_s.fill = fill_banner_sub
    cell_s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, col_count + 1):
        ws.cell(row=2, column=c).fill = fill_banner_sub

    # 3. Boşluk
    ws.row_dimensions[3].height = 6

    # 4. Tablo Kolon Başlıkları
    ws.row_dimensions[4].height = 28
    for col_idx, col_name in enumerate(columns, start=1):
        cell_h = ws.cell(row=4, column=col_idx, value=str(col_name).upper())
        cell_h.font = font_header
        cell_h.fill = fill_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = header_border

    # Sütun Meta Bilgilerini Dinamik Tespit Et
    cat_col = next((c for c in columns if any(k in c.lower() for k in ["kategori", "özelkod6", "grup"])), None)
    
    numeric_cols = set()
    currency_cols = set()
    fark_cols = set()
    percent_cols = set()
    unit_price_cols = set()

    for col_name in columns:
        cl = str(col_name).lower()
        if any(w in cl for w in ["birim fiyat", "alış fiyat", "satış fiyat", "fiyatı", "birim maliyet", "fiyat (tl)", "fiyat(tl)"]):
            unit_price_cols.add(col_name)
        if any(w in cl for w in ["ciro", "tutar", "fiyat", "bedel", "değer", "(tl)", "maliyet"]):
            currency_cols.add(col_name)
        if "fark" in cl:
            fark_cols.add(col_name)
        if any(w in cl for w in ["oran", "yüzde", "%"]):
            percent_cols.add(col_name)

        # Veri örneği üzerinden sayısal kontrol
        num_valid = 0
        for row in data[:30]:
            v = row.get(col_name)
            if v is not None and v != "-" and v != "":
                try:
                    float(v)
                    num_valid += 1
                except (ValueError, TypeError):
                    pass
        if num_valid >= max(1, min(len(data), 5)):
            numeric_cols.add(col_name)

    # Kategoriye göre gruplama yapılabilir mi?
    use_grouping = False
    cat_groups = OrderedDict()
    if cat_col:
        for r in data:
            c_val = str(r.get(cat_col) or "").strip()
            if not c_val or c_val.lower() == "none":
                c_val = "DİĞER"
            if c_val not in cat_groups:
                cat_groups[c_val] = []
            cat_groups[c_val].append(r)
        if len(cat_groups) > 1 or (len(cat_groups) == 1 and list(cat_groups.keys())[0] != "DİĞER"):
            use_grouping = True

    current_row = 5

    def write_row_data(item, row_num, is_alt_row):
        ws.row_dimensions[row_num].height = 20
        base_fill = fill_alt if is_alt_row else fill_white
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = thin_border
            cell.fill = base_fill
            val = item.get(col_name)

            if val is None or val == "" or val == "-":
                cell.value = "-"
                cell.font = font_data
                cell.alignment = Alignment(horizontal="center", vertical="center")
                continue

            if col_name in numeric_cols:
                try:
                    n_val = float(val)
                    cell.value = n_val
                    if col_name in percent_cols:
                        cell.number_format = '+0.0%;-0.0%;0.0%'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif col_name in currency_cols:
                        cell.number_format = '₺ #,##0.00;[Red]-₺ #,##0.00;"-"'
                        cell.font = font_data_bold
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif col_name in fark_cols:
                        cell.number_format = '+#,##0.00;-#,##0.00;0.00'
                        if n_val < 0:
                            cell.font = font_red
                            cell.fill = fill_red_soft
                        elif n_val > 0:
                            cell.font = font_green
                            cell.fill = fill_green_soft
                        else:
                            cell.font = font_data
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.number_format = '#,##0.00' if (n_val % 1 != 0) else '#,##0'
                        cell.font = font_data
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    continue
                except (ValueError, TypeError):
                    pass

            cell.value = str(val)
            cell.font = font_data
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    if use_grouping:
        for cat_name, items in cat_groups.items():
            # A. Kategori başlık şeridi
            ws.row_dimensions[current_row].height = 24
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_count)
            cell_cat = ws.cell(row=current_row, column=1, value=f"📁  {cat_name.upper()}  ({len(items)} Kalem Ürün)")
            cell_cat.font = font_cat
            cell_cat.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for c in range(1, col_count + 1):
                ws.cell(row=current_row, column=c).fill = fill_cat
                ws.cell(row=current_row, column=c).border = cat_border
            current_row += 1

            # B. Kalem satırları
            for idx, item in enumerate(items):
                write_row_data(item, current_row, idx % 2 == 1)
                current_row += 1

            # C. Kategori ara toplam satırı
            ws.row_dimensions[current_row].height = 22
            label_span = 2 if col_count > 3 else 1
            if label_span > 1:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=label_span)
            cell_st = ws.cell(row=current_row, column=1, value=f"📊 {cat_name.upper()} TOPLAMI ({len(items)} Kalem)")
            cell_st.font = font_subtotal
            cell_st.alignment = Alignment(horizontal="left", vertical="center", indent=1)

            for c in range(1, col_count + 1):
                ws.cell(row=current_row, column=c).fill = fill_subtotal
                ws.cell(row=current_row, column=c).border = subtotal_border

            for col_idx, col_name in enumerate(columns, start=1):
                if col_idx <= label_span:
                    continue
                cell_val = ws.cell(row=current_row, column=col_idx)
                if col_name in unit_price_cols:
                    cell_val.value = "-"
                    cell_val.font = font_subtotal
                    cell_val.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name in numeric_cols:
                    tot = 0.0
                    for it in items:
                        try:
                            tot += float(it.get(col_name, 0) or 0)
                        except (ValueError, TypeError):
                            pass
                    cell_val.value = tot
                    cell_val.font = font_subtotal
                    cell_val.alignment = Alignment(horizontal="right", vertical="center")
                    if col_name in currency_cols:
                        cell_val.number_format = '₺ #,##0.00;[Red]-₺ #,##0.00;"-"'
                    elif col_name in fark_cols:
                        cell_val.number_format = '+#,##0.00;-#,##0.00;0.00'
                        if tot < 0:
                            cell_val.font = font_red
                            cell_val.fill = fill_red_soft
                        elif tot > 0:
                            cell_val.font = font_green
                            cell_val.fill = fill_green_soft
                    else:
                        cell_val.number_format = '#,##0.00' if (tot % 1 != 0) else '#,##0'
                else:
                    cell_val.value = "-"
                    cell_val.font = font_subtotal
                    cell_val.alignment = Alignment(horizontal="center", vertical="center")

            current_row += 1
            ws.row_dimensions[current_row].height = 4
            current_row += 1
    else:
        for idx, item in enumerate(data):
            write_row_data(item, current_row, idx % 2 == 1)
            current_row += 1

    # D. GENEL TOPLAM SATIRI
    ws.row_dimensions[current_row].height = 26
    label_span = 2 if col_count > 3 else 1
    if label_span > 1:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=label_span)
    cell_gt = ws.cell(row=current_row, column=1, value=f"👑 GENEL TOPLAM ({len(data)} Kalem)")
    cell_gt.font = font_grand
    cell_gt.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for c in range(1, col_count + 1):
        ws.cell(row=current_row, column=c).fill = fill_grand
        ws.cell(row=current_row, column=c).border = grand_border

    for col_idx, col_name in enumerate(columns, start=1):
        if col_idx <= label_span:
            continue
        cell_val = ws.cell(row=current_row, column=col_idx)
        if col_name in unit_price_cols:
            cell_val.value = "-"
            cell_val.font = font_grand
            cell_val.alignment = Alignment(horizontal="center", vertical="center")
        elif col_name in numeric_cols:
            tot = 0.0
            for it in data:
                try:
                    tot += float(it.get(col_name, 0) or 0)
                except (ValueError, TypeError):
                    pass
            cell_val.value = tot
            cell_val.font = font_grand
            cell_val.alignment = Alignment(horizontal="right", vertical="center")
            if col_name in currency_cols:
                cell_val.number_format = '₺ #,##0.00;[Red]-₺ #,##0.00;"-"'
            elif col_name in fark_cols:
                cell_val.number_format = '+#,##0.00;-#,##0.00;0.00'
            else:
                cell_val.number_format = '#,##0.00' if (tot % 1 != 0) else '#,##0'
        else:
            cell_val.value = "-"
            cell_val.font = font_grand
            cell_val.alignment = Alignment(horizontal="center", vertical="center")

    # Kolon genişlikleri (Auto-fit)
    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for r_idx in range(4, min(current_row + 1, 60)):
            v_str = str(ws.cell(row=r_idx, column=col_idx).value or "")
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(13, min(max_len + 4, 45))

    ws.freeze_panes = "A5"


@router.post("/excel")
async def export_styled_excel(payload: Dict[str, Any] = Body(...)):
    """
    Ekranda görünen raporun tam renk, font, format, kategori ve başlık dizaynıyla
    birebir eşleşen profesyonel ve düzenlenebilir bir .xlsx Excel dosyası üretir.
    - AI Studio: Dinamik sütun algılama, akıllı kategori gruplama, ara toplam ve genel toplam.
    - Mutabakat: Standart şube sayım mutabakatı (özet / gelişmiş).
    - Standart modüller: Temiz kurumsal tablo düzeni.
    """
    title = payload.get("title", "Kurumsal Rapor")
    subtitle = payload.get("subtitle", "Canlı Veritabanı Verisi")
    data = payload.get("data", [])
    module_name = (payload.get("module") or "").strip().lower()
    view_name = (payload.get("view") or "").strip().lower()

    if not data or len(data) == 0:
        raise HTTPException(status_code=400, detail="Dışa aktarılacak veri bulunamadı.")

    # Modül Tespiti
    is_studio = (module_name == "ai_studio")
    is_mutabakat = (
        not is_studio and (
            module_name == "mutabakat" or
            view_name in ("ozet", "gelismis") or
            ("Olması Gereken" in data[0] and "Kategori" in data[0] and "Önceki Sayım" in data[0])
        )
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")

    # =========================================================================
    # 🌟 1. AI STUDIO ÖZEL ANALİZ RAPORLARI
    # =========================================================================
    if is_studio:
        ws.title = "AI Raporu"
        build_ai_studio_sheet(ws, title, subtitle, data, now_str)

    # =========================================================================
    # 🌟 2. STOK SAYIM MUTABAKATI & DENETİM (ÖZET / GELİŞMİŞ)
    # =========================================================================
    elif is_mutabakat:
        ws.title = "Sayım Mutabakatı"
        build_mutabakat_sheet(ws, title, subtitle, data, now_str)

    # =========================================================================
    # 🌟 3. STANDART TABLO RAPORLARI (SATIŞ, CİRO, FİNANS VB.)
    # =========================================================================
    else:
        ws.title = "Rapor"
        build_standard_sheet(ws, title, subtitle, data, now_str)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # Windows dosya yolu sınırları (MAX_PATH) için temiz ve kısa dosya adı oluştur
    clean_fn = "".join([c if c.isalnum() or c in (' ', '_', '-') else '_' for c in title]).strip().replace(' ', '_')
    while '__' in clean_fn:
        clean_fn = clean_fn.replace('__', '_')
    clean_fn = clean_fn.strip('_')
    if len(clean_fn) > 40:
        clean_fn = clean_fn[:40].rstrip('_')
    ascii_fn = "".join([c for c in clean_fn if ord(c) < 128]) or "Rapor"
    date_stamp = datetime.now().strftime("%Y%m%d_%H%M")
    full_filename = f"{clean_fn}_{date_stamp}.xlsx"
    ascii_filename = f"{ascii_fn}_{date_stamp}.xlsx"
    encoded_fn = urllib.parse.quote(full_filename)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{ascii_filename}"; filename*=UTF-8\'\'{encoded_fn}'}
    )


@router.post("/multi-branch-mutabakat")
async def export_multi_branch_mutabakat(payload: Dict[str, Any] = Body(...)):
    raise HTTPException(status_code=404, detail="Bu indirme evrensel çekirdekte tanımlı değil.")


@router.get("/executive-package")
@router.post("/executive-package")
async def export_executive_package():
    raise HTTPException(status_code=404, detail="Bu indirme evrensel çekirdekte tanımlı değil.")
